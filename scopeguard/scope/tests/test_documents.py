import io

import docx
import pytest
from django.conf import settings
from django.core.files.uploadedfile import SimpleUploadedFile
from pypdf import PdfWriter
from rest_framework.test import APIClient

from scopeguard.scope import documents

DEMO_PDF = settings.BASE_DIR / "demo" / "demo-contract.pdf"


def docx_bytes() -> bytes:
    document = docx.Document()
    document.add_paragraph("1. Scope")
    document.add_paragraph("Developer will build a five-page website.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Exclusions"
    table.rows[0].cells[1].text = "E-commerce is not included."
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def blank_pdf_bytes() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def test_extracts_pdf_text():
    with DEMO_PDF.open("rb") as file:
        text = documents.extract_text(file, DEMO_PDF.name)

    assert "Online store, e-commerce, or subscription functionality" in text
    assert "Bramblewood Coffee Roasters" in text


def test_extracts_docx_paragraphs_and_tables():
    text = documents.extract_text(io.BytesIO(docx_bytes()), "sow.docx")

    assert "Developer will build a five-page website." in text
    assert "Exclusions | E-commerce is not included." in text


def test_extracts_plain_text_and_tidies_whitespace():
    raw = "﻿Scope:  \r\n\r\n\r\n\r\nFive pages.\r\n".encode()

    assert documents.extract_text(io.BytesIO(raw), "sow.TXT") == "Scope:\n\nFive pages."


def test_falls_back_for_non_utf8_text():
    raw = "Café menu".encode("cp1252")

    assert documents.extract_text(io.BytesIO(raw), "sow.md") == "Café menu"


def test_rejects_pdf_without_text():
    with pytest.raises(documents.DocumentError, match="Scanned PDFs"):
        documents.extract_text(io.BytesIO(blank_pdf_bytes()), "scan.pdf")


def test_rejects_damaged_files():
    with pytest.raises(documents.DocumentError, match="couldn't be read"):
        documents.extract_text(io.BytesIO(b"not really a pdf"), "sow.pdf")
    with pytest.raises(documents.DocumentError, match="couldn't be read"):
        documents.extract_text(io.BytesIO(b"not really a docx"), "sow.docx")


def test_rejects_unsupported_extension():
    with pytest.raises(documents.DocumentError, match="Upload a PDF"):
        documents.extract_text(io.BytesIO(b"x"), "sow.doc")


# ---- API ---------------------------------------------------------------------

URL = "/api/contracts/extract-sow/"


@pytest.fixture
def api(user):
    client = APIClient()
    client.force_authenticate(user)
    return client


@pytest.mark.django_db
def test_extract_sow_requires_authentication():
    upload = SimpleUploadedFile("sow.txt", b"Five pages.")

    response = APIClient().post(URL, {"file": upload}, format="multipart")

    assert response.status_code in {401, 403}


@pytest.mark.django_db
def test_extract_sow_returns_text(api):
    upload = SimpleUploadedFile("sow.docx", docx_bytes())

    response = api.post(URL, {"file": upload}, format="multipart")

    assert response.status_code == 200  # noqa: PLR2004
    assert response.json()["data"]["filename"] == "sow.docx"
    assert "five-page website" in response.json()["data"]["text"]


@pytest.mark.django_db
@pytest.mark.parametrize(
    ("name", "content", "error"),
    [
        ("sow.exe", b"MZ", "Upload a PDF"),
        ("scan.pdf", blank_pdf_bytes(), "Scanned PDFs"),
    ],
)
def test_extract_sow_reports_unusable_files(api, name, content, error):
    upload = SimpleUploadedFile(name, content)

    response = api.post(URL, {"file": upload}, format="multipart")

    assert response.status_code == 400  # noqa: PLR2004
    assert error in response.json()["errors"]["file"][0]


@pytest.mark.django_db
def test_extract_sow_rejects_large_files(api, monkeypatch):
    monkeypatch.setattr(documents, "MAX_UPLOAD_BYTES", 10)
    upload = SimpleUploadedFile("sow.txt", b"x" * 11)

    response = api.post(URL, {"file": upload}, format="multipart")

    assert response.status_code == 400  # noqa: PLR2004
    assert "limited to" in response.json()["errors"]["file"][0]
